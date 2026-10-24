import csv
import io
import json

from openpyxl import load_workbook

from app.tools.comment_collector import export

ROWS = [
    {"peer_id": 500, "kind": "user", "name": "Иван Петров", "username": "ivan", "link": "https://t.me/ivan",
     "comments": 3, "posts": 2, "first_at": "2026-01-01T00:01:00+00:00", "last_at": "2026-01-01T00:06:00+00:00",
     "is_premium": True, "is_bot": False, "is_verified": False, "is_scam": False, "is_deleted": False, "phone": None},
    {"peer_id": -1003333, "kind": "channel", "name": "Some channel", "username": None, "link": None,
     "comments": 1, "posts": 1, "first_at": None, "last_at": None,
     "is_premium": False, "is_bot": False, "is_verified": False, "is_scam": False, "is_deleted": False, "phone": None},
]


def test_csv():
    data = export.to_csv(ROWS).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(data)))
    assert rows[0][0] == "ID" and len(rows) == 3
    assert rows[1][2] == "Иван Петров" and rows[1][3] == "@ivan" and rows[1][9] == "да"


def test_json():
    data = json.loads(export.to_json(ROWS, {"x": 1}))
    assert data["meta"] == {"x": 1} and data["users"][0]["username"] == "ivan"


def test_xlsx():
    wb = load_workbook(io.BytesIO(export.to_xlsx(ROWS)))
    ws = wb.active
    assert ws.max_row == 3
    assert ws.cell(2, 4).value == "@ivan"
    assert ws.cell(2, 5).hyperlink.target == "https://t.me/ivan"
