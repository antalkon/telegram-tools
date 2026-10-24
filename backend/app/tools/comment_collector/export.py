import csv
import io
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS: list[tuple[str, str]] = [
    ("peer_id", "ID"),
    ("kind", "Тип"),
    ("name", "Имя"),
    ("username", "Username"),
    ("link", "Ссылка"),
    ("comments", "Комментариев"),
    ("posts", "Постов"),
    ("first_at", "Первый комментарий"),
    ("last_at", "Последний комментарий"),
    ("is_premium", "Premium"),
    ("is_bot", "Бот"),
    ("is_verified", "Верифицирован"),
    ("is_scam", "Scam/Fake"),
    ("is_deleted", "Удалён"),
    ("phone", "Телефон"),
]

# chats: messages instead of comments, plus the member list
CHAT_COLUMNS: list[tuple[str, str]] = [
    ("peer_id", "ID"), ("kind", "Тип"), ("name", "Имя"), ("username", "Username"), ("link", "Ссылка"),
    ("comments", "Сообщений"), ("posts", "Дней/тем"), ("member", "Участник"), ("role", "Роль"),
    ("first_at", "Первое сообщение"), ("last_at", "Последнее сообщение"),
    ("is_premium", "Premium"), ("is_bot", "Бот"), ("is_verified", "Верифицирован"), ("is_scam", "Scam/Fake"),
    ("is_deleted", "Удалён"), ("phone", "Телефон"),
]


def _cell(row: dict[str, Any], key: str) -> Any:
    v = row.get(key)
    if key == "username" and v:
        return f"@{v}"
    if isinstance(v, bool):
        return "да" if v else ""
    if key in ("first_at", "last_at") and v:
        return v[:16].replace("T", " ")  # 2026-01-01 00:01 (UTC)
    return v


def to_csv(rows: list[dict[str, Any]], columns: list[tuple[str, str]] = COLUMNS) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title for _, title in columns])
    for r in rows:
        w.writerow([_cell(r, k) for k, _ in columns])
    return ("﻿" + buf.getvalue()).encode()  # BOM so Excel opens UTF-8 correctly


def to_json(rows: list[dict[str, Any]], meta: dict[str, Any]) -> bytes:
    return json.dumps({"meta": meta, "users": rows}, ensure_ascii=False, indent=2).encode()


def to_xlsx(rows: list[dict[str, Any]], sheet_title: str = "Комментаторы", columns: list[tuple[str, str]] = COLUMNS) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append([title for _, title in columns])
    link_col = next((i for i, (k, _) in enumerate(columns, start=1) if k == "link"), None)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E293B")
    for r in rows:
        ws.append([_cell(r, k) for k, _ in columns])
        if (link := r.get("link")) and link_col:
            cell = ws.cell(row=ws.max_row, column=link_col)
            cell.hyperlink = link
            cell.font = Font(color="2AABEE", underline="single")
    widths = {"peer_id": 16, "name": 28, "username": 20, "link": 30, "first_at": 22, "last_at": 22}
    for i, (key, title) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, max(len(title) + 2, 10))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
